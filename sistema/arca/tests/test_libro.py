import csv
import io
from datetime import date, datetime
from decimal import Decimal

from django.urls import reverse
from openpyxl import load_workbook

from payment_instructions.models import User
from sistema.arca import servicios as arca
from sistema.arca.models import ComprobanteFiscal, Emisor
from sistema.arca.servicios import TITULOS_LIBRO
from sistema.arca.tests.fabricas import ArcaFalso, cobro_facturable, emisor_prueba, factura_prueba
from sistema.core.tests import SistemaTestCase
from sistema.core.tiempo import ARG

OCTUBRE = date(2026, 10, 1)


class LibroTestBase(SistemaTestCase):
    def setUp(self):
        super().setUp()
        self.emisor = emisor_prueba()
        self.admin = User.objects.create_user(
            "adm", email="adm@example.com", password="x", role="administrator"
        )
        self.client.force_login(self.admin)

    def emitido(self, fecha, numero, *, tipo="factura", emisor=None, **campos):
        nc = tipo == "nota_credito"
        datos = dict(
            emisor=emisor or self.emisor,
            tipo=tipo,
            estado="emitido",
            cbte_tipo=8 if nc else 6,
            punto_venta=3,
            numero=numero,
            fecha=fecha,
            neto="53719.01",
            iva="11280.99",
            total="65000.00",
            cae="75123456789012",
            cae_vence=date(2026, 11, 10),
            receptor_nombre="Paciente Uno",
            doc_tipo=96,
            doc_numero="30111222",
        )
        if nc:
            origen = factura_prueba(cobro_facturable(), emisor=emisor or self.emisor)
            datos["cobro"], datos["asociado"] = origen.cobro, origen
        else:
            datos["cobro"] = cobro_facturable()
        return ComprobanteFiscal.objects.create(
            **(datos | campos), concepto="servicios", descripcion="Servicios odontológicos"
        )

    def octubre(self):
        for n in (1, 2, 3):
            self.emitido(date(2026, 10, n), n)
        self.emitido(date(2026, 10, 4), 4, tipo="nota_credito")


class LibroIvaTest(LibroTestBase):
    def test_tres_facturas_y_una_nota_de_credito_dan_4_filas_y_la_nota_resta(self):
        self.octubre()
        filas, totales = arca.libro_iva(self.emisor, OCTUBRE)
        self.assertEqual(len(filas), 4)
        self.assertEqual(totales, (Decimal("107438.02"), Decimal("22561.98"), Decimal("130000.00")))
        self.assertEqual(filas[3][1], "Nota de crédito B")
        self.assertEqual(filas[3][8], Decimal("-65000.00"))
        self.assertEqual(filas[0][1:6], ["Factura B", 3, 1, "Paciente Uno", "DNI 30111222"])

    def test_la_factura_del_31_entra_y_la_del_1_no(self):
        self.emitido(date(2026, 10, 31), 1)
        self.emitido(date(2026, 11, 1), 2)
        self.emitido(date(2026, 9, 30), 3)
        filas, _ = arca.libro_iva(self.emisor, OCTUBRE)
        self.assertEqual([f[3] for f in filas], [1])

    def test_diciembre_no_incluye_enero(self):
        self.emitido(date(2026, 12, 31), 1)
        self.emitido(date(2027, 1, 1), 2)
        filas, _ = arca.libro_iva(self.emisor, date(2026, 12, 1))
        self.assertEqual([f[3] for f in filas], [1])

    def test_el_corte_de_mes_usa_la_hora_argentina_al_emitir(self):
        # 31/10 23:59:59 ARG es 01/11 02:59:59 UTC; 01/11 00:00:00 ARG ya es noviembre.
        fake = ArcaFalso(self)
        fake.ultimo = 0
        for momento in (
            datetime(2026, 10, 31, 23, 59, 59, tzinfo=ARG),
            datetime(2026, 11, 1, 0, 0, 0, tzinfo=ARG),
        ):
            pendiente = factura_prueba(cobro_facturable(), emisor=self.emisor)
            self.assertEqual(arca.emitir(pendiente.pk, ahora=momento), "emitido")
        octubre, _ = arca.libro_iva(self.emisor, OCTUBRE)
        noviembre, _ = arca.libro_iva(self.emisor, date(2026, 11, 1))
        self.assertEqual([f[0] for f in octubre], [date(2026, 10, 31)])
        self.assertEqual([f[0] for f in noviembre], [date(2026, 11, 1)])

    def test_pendiente_error_y_descartado_no_entran(self):
        for estado, extra in (
            ("pendiente", {}),
            ("error", {"motivo": "10016: dato inválido"}),
            ("descartado", {"vigente": False}),
        ):
            factura_prueba(
                cobro_facturable(), emisor=self.emisor, estado=estado, fecha=OCTUBRE, **extra
            )
        self.assertEqual(arca.libro_iva(self.emisor, OCTUBRE)[0], [])

    def test_otro_emisor_no_entra(self):
        otro = emisor_prueba(cuit="20000000002", razon_social="Otra S.A.")
        self.emitido(OCTUBRE, 1, emisor=otro)
        self.assertEqual(arca.libro_iva(self.emisor, OCTUBRE)[0], [])

    def test_mes_sin_comprobantes_da_cero_filas_y_totales_en_cero(self):
        filas, totales = arca.libro_iva(self.emisor, OCTUBRE)
        self.assertEqual(filas, [])
        self.assertEqual(totales, (Decimal("0.00"),) * 3)

    def test_una_factura_con_su_nota_de_credito_suma_cero(self):
        self.emitido(OCTUBRE, 1)
        factura = ComprobanteFiscal.objects.get(numero=1)
        self.emitido(OCTUBRE, 2, tipo="nota_credito", asociado=factura, cobro=factura.cobro)
        self.assertEqual(arca.libro_iva(self.emisor, OCTUBRE)[1], (Decimal("0.00"),) * 3)

    def test_orden_por_fecha_tipo_y_numero(self):
        self.emitido(date(2026, 10, 5), 9)
        self.emitido(date(2026, 10, 5), 7)
        self.emitido(date(2026, 10, 2), 8)
        filas, _ = arca.libro_iva(self.emisor, OCTUBRE)
        self.assertEqual([f[3] for f in filas], [8, 7, 9])


class LibroPantallaTest(LibroTestBase):
    def url(self, nombre="arca:libro", **query):
        return reverse(nombre) + "?" + "&".join(f"{k}={v}" for k, v in query.items())

    def test_la_pantalla_muestra_las_filas_y_los_totales(self):
        self.octubre()
        r = self.client.get(self.url(emisor=self.emisor.pk, mes="2026-10"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.context["filas"]), 4)
        self.assertContains(r, "$130.000,00")
        self.assertContains(r, "−$65.000,00")

    def test_emisor_inválido_no_muestra_otro(self):
        self.emitido(OCTUBRE, 1)
        r = self.client.get(self.url(emisor=99999, mes="2026-10"))
        self.assertEqual(len(r.context["filas"]), 0)
        self.assertContains(r, "Elegí un emisor de la lista.")

    def test_solo_el_administrador_entra(self):
        operador = User.objects.create_user(
            "op", email="op@example.com", password="x", role="operator"
        )
        self.client.force_login(operador)
        for nombre in ("arca:libro", "arca:libro_descargar"):
            self.assertEqual(self.client.get(self.url(nombre, formato="csv")).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url()).status_code, 302)

    def test_descarga_csv_con_titulos_filas_y_totales(self):
        self.octubre()
        r = self.client.get(
            self.url("arca:libro_descargar", emisor=self.emisor.pk, mes="2026-10", formato="csv")
        )
        self.assertEqual(r.status_code, 200)
        filas = list(
            csv.reader(
                io.StringIO(b"".join(r.streaming_content).decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(filas[0], TITULOS_LIBRO)
        self.assertEqual(len(filas), 1 + 4 + 1)
        self.assertEqual(filas[-1][1], "Totales")
        self.assertEqual(filas[-1][8], "130000,00")

    def test_descarga_xlsx_con_total_numerico_de_2_decimales(self):
        self.octubre()
        r = self.client.get(
            self.url("arca:libro_descargar", emisor=self.emisor.pk, mes="2026-10", formato="xlsx")
        )
        self.assertEqual(r.status_code, 200)
        hoja = load_workbook(io.BytesIO(r.content)).active
        filas = list(hoja.iter_rows())
        self.assertEqual(len(filas), 6)
        total = filas[-1][8]
        self.assertEqual(total.value, 130000)
        self.assertEqual(total.number_format, "#,##0.00")
        self.assertEqual(filas[4][8].value, -65000)

    def test_receptor_que_empieza_con_igual_sale_como_texto(self):
        self.emitido(OCTUBRE, 1, receptor_nombre="=1+1")
        csv_ = self.client.get(
            self.url("arca:libro_descargar", emisor=self.emisor.pk, mes="2026-10", formato="csv")
        )
        self.assertIn("'=1+1", b"".join(csv_.streaming_content).decode("utf-8-sig"))
        xlsx = self.client.get(
            self.url("arca:libro_descargar", emisor=self.emisor.pk, mes="2026-10", formato="xlsx")
        )
        celda = list(load_workbook(io.BytesIO(xlsx.content)).active.iter_rows())[1][4]
        self.assertEqual(celda.data_type, "s")

    def test_formato_inválido_o_sin_emisor_da_404(self):
        self.assertEqual(
            self.client.get(self.url("arca:libro_descargar", formato="pdf")).status_code, 404
        )
        Emisor.objects.all()._raw_delete("default")  # delete() está bloqueado
        self.assertEqual(
            self.client.get(self.url("arca:libro_descargar", formato="csv")).status_code, 404
        )
