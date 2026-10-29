import zipfile
from datetime import datetime
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.messages import get_messages
from django.urls import reverse
from openpyxl import load_workbook

from sistema.cobros.models import Cobro
from sistema.core.tests import SistemaTestCase
from sistema.core.tiempo import ARG
from sistema.destinatarios.models import Cuenta
from sistema.reportes import servicios

from . import DatosMartinO

TEXTO = (
    "1. 31/10/2026 14:05 - $70.000,00 - Destinatario Ejemplo — alias.uno - Juan Paciente\n"
    "2. 31/10/2026 15:30 - $25.000,00 - Destinatario Ejemplo — alias.dos - Ana Paciente\n\n"
    "Total sin cancelados $95.000,00"
)


class ElegidosTest(DatosMartinO, SistemaTestCase):
    def setUp(self):
        super().setUp()
        self.url = reverse("reportes:cobros_elegidos")
        self.client.force_login(self.admin)

    def post(self, cobros, **datos):
        return self.client.post(self.url, {"pk": [c.pk for c in cobros], **datos})

    def test_texto_portapapeles_ordena_por_fecha_con_total(self):
        b = self.cobro(
            self.galicia, "25000.00", datetime(2026, 10, 31, 15, 30, tzinfo=ARG),
            nombre_paciente="Ana Paciente",
        )  # fmt: skip
        a = self.cobro(
            self.tincho, "70000.00", datetime(2026, 10, 31, 14, 5, tzinfo=ARG),
            nombre_paciente="Juan Paciente",
        )  # fmt: skip
        qs = Cobro.objects.filter(pk__in=[b.pk, a.pk])
        self.assertEqual(servicios.texto_portapapeles(qs), TEXTO)
        r = self.post([b, a], accion="copiar")
        self.assertEqual(r.status_code, 200)
        self.assertContains(
            r, "1. 31/10/2026 14:05 - $70.000,00 - Destinatario Ejemplo — alias.uno"
        )
        self.assertContains(r, "Total sin cancelados $95.000,00")

    def test_texto_portapapeles_vacio_sin_cobros(self):
        self.assertEqual(servicios.texto_portapapeles(Cobro.objects.none()), "")

    def test_portapapeles_pos_usa_titular_y_suma_centavos(self):
        a = self.cobro(self.mp, "0.01")
        b = self.cobro(self.mp, "0.01")
        texto = servicios.texto_portapapeles(Cobro.objects.filter(pk__in=[a.pk, b.pk]))
        self.assertIn("- Destinatario Ejemplo — Martin MP -", texto)
        self.assertTrue(texto.endswith("Total sin cancelados $0,02"))

    def subir(self, cobro, nombre):
        cobro.comprobante.save(nombre, SimpleUploadedFile(nombre, b"datos"))

    def test_zip_con_4_comprobantes_y_un_efectivo(self):
        efectivo = Cuenta.objects.create(destinatario=self.martin, tipo="efectivo", titular="Caja")
        cobros = [self.cobro(self.tincho, "100.00") for _ in range(4)]
        for c in cobros:
            self.subir(c, "x.pdf")
        cobros.append(self.cobro(efectivo, "50.00", metodo="efectivo"))
        r = self.post(cobros, accion="comprobantes")
        self.assertEqual(r["Content-Type"], "application/zip")
        self.assertIn("comprobantes-", r["Content-Disposition"])
        nombres = zipfile.ZipFile(BytesIO(b"".join(r.streaming_content))).namelist()
        self.assertCountEqual(
            nombres, [f"{c.pk}_alias.uno_15_10_2026_12_00.pdf" for c in cobros[:4]]
        )

    def test_zip_sin_alias_y_archivo_faltante(self):
        c = self.cobro(self.mp, "100.00")  # comprobante de la fábrica: no existe en el storage
        d = self.cobro(self.mp, "100.00")
        self.subir(d, "y.png")
        qs = Cobro.objects.filter(pk__in=[c.pk, d.pk])
        nombres = zipfile.ZipFile(servicios.zip_comprobantes(qs)).namelist()
        self.assertEqual(nombres, [f"{d.pk}_sin_alias_15_10_2026_12_00.png"])

    def test_solo_efectivo_o_archivo_faltante_redirige_con_aviso(self):
        efectivo = Cuenta.objects.create(destinatario=self.martin, tipo="efectivo", titular="Caja")
        c = self.cobro(efectivo, "50.00", metodo="efectivo")
        r = self.post([c], accion="comprobantes")
        self.assertRedirects(r, reverse("cobros:lista"), fetch_redirect_response=False)
        self.assertEqual(
            [str(m) for m in get_messages(r.wsgi_request)],
            ["No se encontraron comprobantes en los cobros elegidos."],
        )
        faltante = self.cobro(self.tincho, "100.00")
        self.assertIsNone(servicios.zip_comprobantes(Cobro.objects.filter(pk=faltante.pk)))

    def test_descargar_xlsx_con_2_elegidos(self):
        a, b, _ = (self.cobro(self.tincho, "100.00") for _ in range(3))
        r = self.post([a, b], accion="descargar", formato="xlsx", por="cobro")
        self.assertEqual(load_workbook(BytesIO(r.content)).active.max_row, 3)

    def test_descargar_con_formato_invalido_da_400(self):
        a = self.cobro(self.tincho, "100.00")
        self.assertEqual(
            self.post([a], accion="descargar", formato="pdf", por="cobro").status_code, 400
        )
        self.assertEqual(self.post([a], accion="otra").status_code, 400)
