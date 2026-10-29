from datetime import datetime
from io import BytesIO
from unittest import mock

from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from payment_instructions.models import User
from sistema.auditoria.models import Evento
from sistema.cobros.models import ItemCobro
from sistema.core.tests import SistemaTestCase
from sistema.core.tiempo import ARG

from . import D, EN_OCT, DatosMartinO


class DescargasCobrosTest(DatosMartinO, SistemaTestCase):
    def setUp(self):
        super().setUp()
        self.url = reverse("reportes:cobros_descargar")
        self.client.force_login(self.admin)

    def get(self, **q):
        return self.client.get(self.url, q)

    def test_csv_por_cobro_filtra_por_metodo_con_bom_y_acentos(self):
        for _ in range(3):
            self.cobro(self.tincho, "100.00", nombre_paciente="Núñez Peña")
        efectivo = self.cuenta_efectivo()
        for _ in range(2):
            self.cobro(efectivo, "50.00", metodo="efectivo")
        r = self.get(metodo="transferencia", formato="csv", por="cobro")
        self.assertEqual(r.status_code, 200)
        contenido = b"".join(r.streaming_content)
        self.assertTrue(contenido.startswith("\ufeff".encode()))
        lineas = contenido.decode("utf-8-sig").strip().splitlines()
        self.assertEqual(len(lineas), 4)
        self.assertIn("Núñez Peña", lineas[1])
        self.assertEqual(lineas[0].count(";"), 21)

    def cuenta_efectivo(self):
        from sistema.destinatarios.models import Cuenta

        return Cuenta.objects.create(destinatario=self.martin, tipo="efectivo", titular="Caja")

    def test_xlsx_por_item_numeros_y_fecha_arg(self):
        c = self.cobro(self.tincho, "70000.00", fecha=datetime(2026, 10, 31, 23, 59, tzinfo=ARG))
        ItemCobro.objects.create(cobro=c, numero=1, id_prestacion=1, monto=D("70000.00"))
        r = self.get(formato="xlsx", por="item")
        hoja = load_workbook(BytesIO(r.content)).active
        self.assertEqual(hoja.max_row, 2)
        self.assertEqual(hoja["B2"].value, datetime(2026, 10, 31, 23, 59))
        self.assertEqual(hoja["E2"].value, 70000)
        self.assertEqual(hoja["E2"].number_format, "#,##0.00")

    def test_sin_resultados_solo_titulos(self):
        r = self.get(formato="csv", por="cobro", metodo="qr")
        self.assertEqual(
            len(b"".join(r.streaming_content).decode("utf-8-sig").strip().splitlines()), 1
        )

    def test_nombre_de_archivo_con_fecha_argentina(self):
        ahora = datetime(2026, 11, 5, 10, 0, tzinfo=ARG)
        with mock.patch("sistema.reportes.descarga.timezone", wraps=timezone) as tz:
            tz.now.return_value = ahora
            a = self.get(formato="xlsx", por="cobro")
            b = self.get(formato="csv", por="item")
        self.assertIn('filename="cobros-05-11-2026.xlsx"', a["Content-Disposition"])
        self.assertIn('filename="cobros-por-item-05-11-2026.csv"', b["Content-Disposition"])

    def test_parametros_invalidos_dan_400(self):
        self.assertEqual(self.get(formato="csv", por="otro").status_code, 400)
        self.assertEqual(self.get(formato="pdf", por="cobro").status_code, 400)
        self.assertEqual(self.get(por="cobro").status_code, 400)
        self.assertEqual(self.get(formato="csv").status_code, 400)

    def test_cancelados_salen_si_el_filtro_los_incluye(self):
        self.cobro(self.tincho, "100.00", estado="cancelado")
        self.cobro(self.tincho, "100.00")
        todos = b"".join(self.get(formato="csv", por="cobro").streaming_content).decode("utf-8-sig")
        solo = b"".join(
            self.get(formato="csv", por="cobro", estado="cancelado").streaming_content
        ).decode("utf-8-sig")
        self.assertEqual(len(todos.strip().splitlines()), 3)
        self.assertEqual(len(solo.strip().splitlines()), 2)
        self.assertIn("Cancelado", solo)

    def test_operador_y_profesional_reciben_403(self):
        op = User.objects.create_user("ana", email="a@example.com", role="operator")
        pro = User.objects.create_user("pro", email="p@example.com", role="professional")
        for u in (op, pro):
            self.client.force_login(u)
            self.assertEqual(self.get(formato="csv", por="cobro").status_code, 403)

    def test_get_no_crea_eventos(self):
        self.cobro(self.tincho, "100.00")
        antes = Evento.objects.count()
        self.get(formato="csv", por="item")
        self.assertEqual(Evento.objects.count(), antes)

    def test_consultas_no_dependen_de_la_cantidad_de_cobros(self):
        def cuenta_consultas():
            with CaptureQueriesContext(connection) as q:
                self.get(formato="csv", por="item")
            return len(q)

        for n in range(2):
            c = self.cobro(self.tincho, "100.00", fecha=EN_OCT)
            ItemCobro.objects.create(cobro=c, numero=1, id_prestacion=n, monto=D("100.00"))
        con_2 = cuenta_consultas()
        for n in range(18):
            c = self.cobro(self.tincho, "100.00", fecha=EN_OCT)
            ItemCobro.objects.create(cobro=c, numero=1, id_prestacion=n, monto=D("100.00"))
        self.assertEqual(cuenta_consultas(), con_2)

    def test_total_recibido_del_reporte_es_igual_a_la_suma_de_la_descarga(self):
        # SC-001: sin pagos complementarios, reporte == suma de la columna Monto
        from sistema.reportes.servicios import reporte_destinatario

        self.cargar_octubre()
        r = self.get(
            formato="csv",
            por="cobro",
            destinatario=self.martin.pk,
            desde="2026-10-01",
            hasta="2026-10-31",
        )
        filas = b"".join(r.streaming_content).decode("utf-8-sig").strip().splitlines()
        cols = filas[0].split(";")
        i = cols.index("Monto")
        suma = sum(D(f.split(";")[i].replace(",", ".")) for f in filas[1:])
        self.assertEqual(
            suma, reporte_destinatario(self.martin, EN_OCT.date().replace(day=1)).recibido
        )


class ZipComprobantesTest(DatosMartinO, SistemaTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)

    def zip(self, cobro):
        import zipfile

        from sistema.cobros.models import Cobro
        from sistema.reportes import servicios

        datos = servicios.zip_comprobantes(Cobro.objects.filter(pk=cobro.pk))
        return zipfile.ZipFile(datos).namelist()

    def test_zip_trae_todos_los_comprobantes_del_cobro(self):
        from django.core.files.base import ContentFile

        from sistema.cobros.models import ComprobanteCobro

        c1 = self.cobro(self.tincho, "100.00", estado="pendiente")
        c1.comprobante.save("a.jpg", ContentFile(b"x"), save=True)
        solo = self.zip(c1)
        self.assertEqual(len(solo), 1)
        for _ in range(2):
            otro = ComprobanteCobro(cobro=c1, creado_por=self.admin)
            otro.archivo.save("b.jpg", ContentFile(b"y"), save=False)
            otro.save()
        nombres = self.zip(c1)
        self.assertEqual(len(nombres), 3)
        self.assertTrue(nombres[1].endswith("_2.jpg") and nombres[2].endswith("_3.jpg"))

    def test_csv_enlaza_todos_los_comprobantes_uno_por_linea(self):
        import csv
        from io import StringIO

        from django.core.files.base import ContentFile

        from sistema.cobros.models import ComprobanteCobro

        c1 = self.cobro(self.tincho, "100.00", estado="pendiente")
        c1.comprobante.save("a.jpg", ContentFile(b"x"), save=True)
        otros = []
        for _ in range(2):
            otro = ComprobanteCobro(cobro=c1, creado_por=self.admin)
            otro.archivo.save("b.jpg", ContentFile(b"y"), save=False)
            otro.save()
            otros.append(otro)
        r = self.client.get(
            reverse("reportes:cobros_descargar"), {"formato": "csv", "por": "cobro"}
        )
        filas = list(
            csv.reader(StringIO(b"".join(r.streaming_content).decode("utf-8-sig")), delimiter=";")
        )
        celda = filas[1][filas[0].index("Link al comprobante")]
        enlaces = celda.split("\n")
        self.assertEqual(len(enlaces), 3)
        self.assertTrue(enlaces[0].endswith(reverse("cobros:comprobante", args=[c1.pk])))
        for enlace, otro in zip(enlaces[1:], otros, strict=True):
            self.assertTrue(enlace.endswith(reverse("cobros:comprobante_otro", args=[otro.pk])))

    def test_un_solo_comprobante_sigue_siendo_un_enlace(self):
        import csv
        from io import StringIO

        from django.core.files.base import ContentFile

        c1 = self.cobro(self.tincho, "100.00", estado="pendiente")
        c1.comprobante.save("a.jpg", ContentFile(b"x"), save=True)
        r = self.client.get(
            reverse("reportes:cobros_descargar"), {"formato": "csv", "por": "cobro"}
        )
        filas = list(
            csv.reader(StringIO(b"".join(r.streaming_content).decode("utf-8-sig")), delimiter=";")
        )
        self.assertNotIn("\n", filas[1][filas[0].index("Link al comprobante")])
