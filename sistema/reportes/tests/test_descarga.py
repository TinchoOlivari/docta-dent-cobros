from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from sistema.core.tests import SistemaTestCase
from sistema.core.tiempo import ARG
from sistema.reportes.descarga import respuesta

D = Decimal
TITULOS = ["Nombre", "Monto", "Fecha"]
FILA = ["Núñez Peña", D("1234.56"), datetime(2026, 10, 31, 23, 59, 59, tzinfo=ARG)]


class DescargaTest(SistemaTestCase):
    def test_csv_con_bom_punto_y_coma_y_coma_decimal(self):
        r = respuesta("cobros", TITULOS, [FILA], "csv")
        contenido = b"".join(r.streaming_content)
        self.assertTrue(contenido.startswith("\ufeff".encode()))
        self.assertIn("Núñez Peña;1234,56;31/10/2026 23:59", contenido.decode("utf-8-sig"))
        self.assertEqual(r["Content-Type"], "text/csv; charset=utf-8")

    def test_csv_none_sale_vacio(self):
        r = respuesta("cobros", TITULOS, [["a", None, None]], "csv")
        self.assertIn("a;;", b"".join(r.streaming_content).decode("utf-8-sig"))

    def test_xlsx_numeros_y_fechas_reales(self):
        r = respuesta("cobros", TITULOS, [FILA], "xlsx")
        hoja = load_workbook(BytesIO(r.content)).active
        self.assertEqual(hoja["A2"].value, "Núñez Peña")
        self.assertEqual(hoja["B2"].value, 1234.56)
        self.assertEqual(hoja["B2"].number_format, "#,##0.00")
        self.assertEqual(hoja["C2"].value, datetime(2026, 10, 31, 23, 59, 59))
        self.assertEqual(hoja["C2"].number_format, "dd/mm/yyyy hh:mm")

    def test_cero_filas_solo_titulos(self):
        csv_ = b"".join(respuesta("cobros", TITULOS, [], "csv").streaming_content).decode(
            "utf-8-sig"
        )
        self.assertEqual(csv_.strip(), "Nombre;Monto;Fecha")
        hoja = load_workbook(BytesIO(respuesta("cobros", TITULOS, [], "xlsx").content)).active
        self.assertEqual(hoja.max_row, 1)
        self.assertEqual(hoja["A1"].value, "Nombre")

    def test_nombre_del_archivo_en_hora_argentina(self):
        r = respuesta("cobros", TITULOS, [], "xlsx", ahora=datetime(2026, 11, 5, 0, 30, tzinfo=ARG))
        self.assertIn('filename="cobros-05-11-2026.xlsx"', r["Content-Disposition"])
        # 5/11 02:00 UTC es 4/11 23:00 en Argentina
        ahora = datetime(2026, 11, 5, 2, 0, tzinfo=timezone.utc)
        r = respuesta("cobros", TITULOS, [], "csv", ahora=ahora)
        self.assertIn('filename="cobros-04-11-2026.csv"', r["Content-Disposition"])

    def test_texto_con_formula_no_corre_como_formula(self):
        fila = ['=HYPERLINK("http://x","a")', "-1+1", "Ana"]
        csv_ = b"".join(respuesta("cobros", TITULOS, [fila], "csv").streaming_content).decode(
            "utf-8-sig"
        )
        self.assertIn('\'=HYPERLINK(""http://x"",""a"")";\'-1+1;Ana', csv_)
        hoja = load_workbook(BytesIO(respuesta("cobros", TITULOS, [fila], "xlsx").content)).active
        self.assertEqual(hoja["A2"].data_type, "s")
        self.assertEqual(hoja["A2"].value, fila[0])
        self.assertEqual(hoja["C2"].value, "Ana")

    def test_formato_invalido(self):
        with self.assertRaises(ValueError):
            respuesta("cobros", TITULOS, [], "pdf")
