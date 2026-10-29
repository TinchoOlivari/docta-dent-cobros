from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from sistema.core.tests import SistemaTestCase
from sistema.destinatarios.servicios import crear_destinatario

from . import DatosMartinO


class DestinatariosDescargaTest(DatosMartinO, SistemaTestCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)
        self.url = reverse("reportes:destinatarios_descargar")
        self.cargar_octubre()

    def test_xlsx_trae_la_fila_de_martin_con_numeros(self):
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "xlsx"})
        filas = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(
            filas[0], ("Mes", "Destinatario", "Esperado", "Recibido", "Falta", "Cuentas")
        )
        self.assertEqual(
            filas[1],
            ("10/2026", "Destinatario Ejemplo", 1150000, 1000000, 150000, "alias.uno, alias.dos"),
        )

    def test_con_falta_deja_afuera_al_destinatario_sin_falta(self):
        sin = crear_destinatario(usuario=None, nombre="Sin Falta", tipo="profesional")
        self.cobro(self.transf(sin, "sin.falta", "1000.00"), "1000.00", destinatario=sin)
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "csv"})
        self.assertIn("Sin Falta", b"".join(r.streaming_content).decode("utf-8-sig"))
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "csv", "con_falta": "1"})
        texto = b"".join(r.streaming_content).decode("utf-8-sig")
        self.assertNotIn("Sin Falta", texto)
        self.assertIn("Destinatario Ejemplo", texto)

    def test_csv_con_punto_y_coma_y_coma_decimal(self):
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "csv"})
        self.assertIn(
            "10/2026;Destinatario Ejemplo;1150000,00;1000000,00;150000,00;",
            b"".join(r.streaming_content).decode("utf-8-sig"),
        )

    def test_formato_invalido_da_400(self):
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "pdf"})
        self.assertEqual(r.status_code, 400)

    def test_sin_formato_da_400(self):
        self.assertEqual(self.client.get(self.url).status_code, 400)

    def test_nombre_de_archivo(self):
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "csv"})
        self.assertTrue(
            r["Content-Disposition"].startswith('attachment; filename="falta-del-mes-2026-10.csv')
        )

    def test_destinatario_sin_cuentas_activas_no_sale(self):
        crear_destinatario(usuario=None, nombre="Sin Cuentas", tipo="clinica")  # 021/FR-026
        r = self.client.get(self.url, {"mes": "2026-10", "formato": "csv"})
        self.assertNotIn("Sin Cuentas", b"".join(r.streaming_content).decode("utf-8-sig"))
