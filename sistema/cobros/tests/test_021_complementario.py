"""021 Fase C: Falta del mes, cuenta por destinatario, descargas, pago migrado (FR-014, 026, 028, 029, 042)."""

import re
from datetime import date
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from sistema.cobros.tests.fabrica import pago_prueba, topes_desde_enero
from sistema.cobros.tests.test_complementario_pantallas import BasePantalla, D
from sistema.destinatarios.servicios import (
    cambiar_estado_destinatario,
    crear_cuenta,
    crear_destinatario,
)

LISTA = reverse("cobros:complementarios")
NUEVO = reverse("cobros:complementario_nuevo")


def texto_csv(r):
    return b"".join(r.streaming_content).decode("utf-8-sig")


class FaltaDelMesTest(BasePantalla):
    def destino(self, recibido):
        d = crear_destinatario(usuario=None, nombre="Destinatario Ejemplo", tipo="profesional")
        c = crear_cuenta(
            usuario=None,
            tipo="transferencia",
            destinatario=d,
            alias="ejemplo.uno",
            tope_mensual=D("100000.00"),
        )
        topes_desde_enero()
        self.cobro(d, c, recibido)
        return d

    def pagina(self):
        # Jimena (de setUp) no recibió nada: su fila queda; las asserts miran el destinatario nuevo
        r = self.client.get(LISTA, {"mes": "2026-10"})
        fila = re.search(
            r"<tr>(?:(?!</tr>).)*Destinatario Ejemplo.*?</tr>", r.content.decode(), re.S
        )
        return fila.group(0)

    def test_recibio_un_centavo_de_mas(self):
        self.destino("100000.01")
        fila = self.pagina()
        self.assertIn("Recibió de más $0,01", fila)
        self.assertNotIn("Registrar pago", fila)

    def test_recibido_justo_en_el_esperado_no_ofrece_registrar_ni_de_mas(self):
        self.destino("100000.00")
        fila = self.pagina()
        self.assertIn("$0,00", fila)
        self.assertNotIn("Registrar pago", fila)
        self.assertNotIn("Recibió de más", fila)

    def test_un_centavo_menos_ofrece_registrar_pago(self):
        self.destino("99999.99")
        fila = self.pagina()
        self.assertIn("$0,01", fila)
        self.assertIn("Registrar pago", fila)
        self.assertNotIn("Recibió de más", fila)

    def test_recibido_cero_ofrece_registrar_pago(self):
        self.destino("0.01")  # el mínimo posible; cero no tiene cobro
        self.assertIn("Registrar pago", self.pagina())

    def test_destinatario_de_senas_sin_cuentas_activas_no_aparece(self):
        crear_destinatario(usuario=None, nombre="Destinatario Señas", tipo="clinica")
        self.assertNotContains(self.client.get(LISTA, {"mes": "2026-10"}), "Destinatario Señas")

    def test_mes_vacio_deshabilita_las_descargas(self):
        r = self.client.get(LISTA, {"mes": ""})
        self.assertEqual(self.client.get(LISTA, {"mes": "", "formato": "csv"}).status_code, 400)


class CuentaPorDestinatarioTest(BasePantalla):
    def setUp(self):
        super().setUp()
        self.d3 = crear_destinatario(
            usuario=None, nombre="Destinatario Ejemplo 3", tipo="profesional"
        )
        for alias in ("ejemplo.tres.a", "ejemplo.tres.b"):
            crear_cuenta(
                usuario=None,
                tipo="transferencia",
                destinatario=self.d3,
                alias=alias,
                tope_mensual=D("1000.00"),
            )

    def opciones(self, **kw):
        html = self.client.get(NUEVO, kw).content.decode()
        select = re.search(r'<select[^>]*name="cuenta".*?</select>', html, re.S).group(0)
        return re.findall(r"<option[^>]*>", select)

    def test_destinatario_archivado_no_ofrece_sus_cuentas(self):
        cambiar_estado_destinatario(self.d3, "archivado", usuario=None)
        ops = self.opciones()
        self.assertFalse(any(f'data-destinatario="{self.d3.pk}"' in o for o in ops))
        self.assertTrue(any(f'data-destinatario="{self.jimena.pk}"' in o for o in ops))


class DescargaPagosTest(BasePantalla):
    def setUp(self):
        super().setUp()
        for m in ("1000.00", "2000.00", "3000.00"):
            pago_prueba(
                destinatario=self.jimena, cuenta=self.cuenta, monto=D(m), creado_por=self.admin
            )
        pago_prueba(  # de otro mes: no sale
            destinatario=self.jimena, cuenta=self.cuenta, monto=D("9.00"), creado_por=self.admin,
            mes=date(2026, 9, 1),
        )  # fmt: skip

    def test_csv_con_nombre_del_mes_y_tres_filas(self):
        r = self.client.get(LISTA, {"mes": "2026-10", "formato": "csv"})
        self.assertIn("pagos-complementarios-2026-10.csv", r["Content-Disposition"])
        lineas = texto_csv(r).strip().splitlines()
        self.assertEqual(len(lineas), 4)
        self.assertIn("1000,00", "".join(lineas))

    def test_xlsx_con_tres_filas_y_numeros(self):
        r = self.client.get(LISTA, {"mes": "2026-10", "formato": "xlsx"})
        self.assertIn("pagos-complementarios-2026-10.xlsx", r["Content-Disposition"])
        filas = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(len(filas), 4)
        self.assertEqual(sorted(f[4] for f in filas[1:]), [1000, 2000, 3000])

    def test_mes_sin_pagos_baja_solo_los_titulos(self):
        r = self.client.get(LISTA, {"mes": "2026-08", "formato": "csv"})
        self.assertEqual(len(texto_csv(r).strip().splitlines()), 1)

    def test_formato_invalido_es_400(self):
        self.assertEqual(
            self.client.get(LISTA, {"mes": "2026-10", "formato": "pdf"}).status_code, 400
        )


class FaltaDelMesDescargaTest(BasePantalla):
    def test_nombre_y_columna_mes(self):
        r = self.client.get(
            reverse("reportes:destinatarios_descargar"), {"mes": "2026-10", "formato": "csv"}
        )
        self.assertIn("falta-del-mes-2026-10.csv", r["Content-Disposition"])
        lineas = texto_csv(r).splitlines()
        self.assertTrue(lineas[0].startswith("Mes;Destinatario"))
        self.assertTrue(lineas[1].startswith("10/2026;Jimena"))
